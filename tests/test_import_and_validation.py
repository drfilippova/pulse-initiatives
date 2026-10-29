"""Тесты импорта Word-отчёта, проверки Excel-строк и выгрузки (T-10, T-12…T-16, T-19)."""

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from pulse import calc, db, loaders
from pulse.export import build_report
from pulse.import_excel import FileReadError, build_template, read_table
from pulse.import_word import parse_docx, parse_number
from pulse.validation import validate

REPORT = loaders.CENTER_REPORT


def test_t12_parse_center_report():
    res = parse_docx(str(REPORT))
    assert len(res.values) == 85
    assert res.n_indicators == 17
    assert res.n_initiatives == 10
    assert sum(1 for k in {(v.initiative_id, v.indicator, v.is_cumulative) for v in res.values} if k[2]) == 3
    assert res.unknown_initiatives == []
    assert not res.problems


def test_t10_question_marks_mean_no_data():
    res = parse_docx(str(REPORT))
    assert "INI-06" in res.no_data
    assert "???" in res.no_data["INI-06"]


@pytest.mark.parametrize("text,expected", [("0,47", 0.47), ("11 300", 11300), ("17", 17), ("3,07", 3.07)])
def test_t13_numbers(text, expected):
    assert parse_number(text) == pytest.approx(expected)


def test_center_report_loads_and_confirms(tmp_db):
    rep = loaders.load_center_report()
    assert rep.added == 85
    v = db.values_frame()
    assert len(v) == 85
    eq = v[(v["year"] == 2026) & (v["plan"] == v["fact"])]
    assert len(eq) == 4 and eq["confirmed"].all()
    s = calc.year_summary(v, 2026, calc.Thresholds(), 2026, db.initiatives_frame())
    assert s.q1_flags == 0
    assert s.lagging == 0
    assert s.counts["high"] == 17          # при данных на 30.06 все показатели «план занижен»
    assert s.with_data == 10


def test_center_report_forecast_above_previous_year(tmp_db):
    """Допущение «данные на 30.06»: прогноз 2026 больше факта 2025 у всех показателей."""
    loaders.load_center_report()
    v = db.values_frame()
    rows = calc.indicator_rows(v, 2026, calc.Thresholds(), 2026)
    prev = v[v["year"] == 2025].set_index("indicator_id")["fact"]
    assert all(r.fc > prev[r.indicator_id] for r in rows.itertuples())


def _indicators():
    return db.indicators_frame()


def test_t14_empty_and_negative(tmp_db):
    loaders.load_center_report()
    ind = _indicators()
    name = ind[ind["initiative_id"] == "INI-09"].iloc[0]["name"]
    df = pd.DataFrame([
        {"Инициатива": "Третий семестр", "Показатель": name, "Год": 2026, "План": 60, "Факт": None, "Данные на месяц": 9},
        {"Инициатива": "Третий семестр", "Показатель": name, "Год": 2025, "План": 45, "Факт": -5, "Данные на месяц": 12},
        {"Инициатива": "Третий семестр", "Показатель": name, "Год": 2024, "План": 32, "Факт": "35", "Данные на месяц": 12},
    ])
    res = validate(df, ind, {}, 2026)
    rules = {(i.rule, i.row) for i in res.errors}
    assert ("V-02", 2) in rules and ("V-04", 3) in rules
    assert len(res.rows) == 1 and res.rows[0]["year"] == 2024


def test_t15_wrong_indicator_and_initiative(tmp_db):
    loaders.load_center_report()
    df = pd.DataFrame([
        {"Инициатива": "Наука рядом", "Показатель": "Количество маршрутов научно-популярного туризма",
         "Год": 2026, "План": 1, "Факт": 1, "Данные на месяц": 6},
        {"Инициатива": "Наука близко", "Показатель": "x", "Год": 2026, "План": 1, "Факт": 1, "Данные на месяц": 6},
        {"Инициатива": "Наука рядом", "Показатель": "x", "Год": 2026, "План": 1, "Факт": 1, "Данные на месяц": 13},
    ])
    res = validate(df, _indicators(), {}, 2026)
    codes = {i.rule for i in res.errors}
    assert {"V-05", "V-06", "V-07"} <= codes
    assert res.rows == []


def test_duplicates_in_file_and_db(tmp_db):
    loaders.load_center_report()
    ind = _indicators()
    name = ind[ind["initiative_id"] == "INI-10"].iloc[0]["name"]
    row = {"Инициатива": "Наука и бизнес", "Показатель": name, "Год": 2026, "План": 36, "Факт": "40",
           "Данные на месяц": 9}
    dup = validate(pd.DataFrame([row, row]), ind, db.existing_keys(), 2026)
    assert {i.rule for i in dup.errors} == {"V-08"} and dup.rows == []
    one = validate(pd.DataFrame([row]), ind, db.existing_keys(), 2026)
    assert [w.rule for w in one.warnings] == ["V-09"]
    rep = loaders.load_validated(one, "test.xlsx", on_duplicate="replace")
    assert rep.replaced == 1
    v = db.values_frame()
    assert v[(v["indicator"] == name) & (v["year"] == 2026) & (v["initiative_id"] == "INI-10")].iloc[0]["fact"] == 40


def test_missing_columns():
    res = validate(pd.DataFrame([{"Инициатива": "Наука рядом"}]), pd.DataFrame(), {}, 2026)
    assert res.file_error and res.errors[0].rule == "V-01"


def test_t18_bad_files():
    with pytest.raises(FileReadError):
        read_table(b"not an excel file", "broken.xlsx")
    with pytest.raises(FileReadError):
        read_table(b"abc", "report.docx")


def test_template_roundtrip(tmp_db):
    loaders.load_center_report()
    ind = _indicators()
    data = build_template(ind, 2026, 6)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Данные", "Справочники", "Инструкция"]
    ws = wb["Данные"]
    ws.append(["Наука как искусство", ind[ind["initiative_id"] == "INI-04"].iloc[0]["name"], 2026, 23, 20, 6])
    buf = BytesIO()
    wb.save(buf)
    df = read_table(buf.getvalue(), "filled.xlsx")
    res = validate(df, ind, db.existing_keys(), 2026)
    assert not res.errors and len(res.rows) == 1


def test_csv_with_semicolon(tmp_db):
    loaders.load_center_report()
    name = _indicators().query("initiative_id == 'INI-12'").iloc[0]["name"]
    text = f"Инициатива;Показатель;Год;План;Факт;Данные на месяц\nРешения и сервисы;{name};2026;23;1 500,5;9\n"
    df = read_table(text.encode("cp1251"), "data.csv")
    res = validate(df, _indicators(), {}, 2026)
    assert not res.errors and res.rows[0]["fact"] == pytest.approx(1500.5)


def test_demo_has_all_statuses(tmp_db):
    loaders.load_demo()
    v = db.values_frame()
    s = calc.year_summary(v, 2026, calc.Thresholds(), 2026, db.initiatives_frame())
    statuses = {i.status for i in s.initiatives}
    assert {"crit", "risk", "ok", "high", "nodata"} <= statuses
    assert s.q1_flags == 1


def test_t19_export(tmp_db):
    loaders.load_center_report()
    v = db.values_frame()
    thr = calc.Thresholds()
    s = calc.year_summary(v, 2026, thr, 2026, db.initiatives_frame())
    data = build_report(s, calc.quality_flags(v, 2026, db.initiatives_frame()), v, thr, "report")
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Сводка", "Детализация", "Флаги качества", "Данные", "Параметры"]
    assert wb["Детализация"].max_row == 1 + 17
    assert wb["Данные"].max_row == 1 + 85


def test_set_current_month(tmp_db):
    loaders.load_center_report()
    db.set_current_month(2026, 9)
    v = db.values_frame()
    assert set(v[v["year"] == 2026]["as_of_month"]) == {9}
    assert db.get_settings()["current_month"] == "9"
