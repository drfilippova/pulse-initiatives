"""Сводный отчёт в Excel: сводка, детализация, флаги качества, исходные данные, параметры."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .calc import FLAG_TITLES, STATUS_LABELS, Thresholds, YearSummary

BLUE = "1E4FC2"
STATUS_FILL = {
    "crit": "FDE3E7", "risk": "FFF0D9", "ok": "DDF3E8", "high": "E2E6FF", "nodata": "ECEFF5", "na": "ECEFF5",
}
STATUS_FONT = {
    "crit": "C4243B", "risk": "A25A00", "ok": "0E7A4F", "high": "2B3FD6", "nodata": "5E6A86", "na": "5E6A86",
}
THIN = Side(style="thin", color="D9E1F2")


def _header(ws, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{ws.max_row + 1}"


def _round(v, digits=1):
    return None if v is None or pd.isna(v) else round(float(v), digits)


def _paint_status(cell, status: str) -> None:
    cell.fill = PatternFill("solid", fgColor=STATUS_FILL.get(status, "FFFFFF"))
    cell.font = Font(bold=True, color=STATUS_FONT.get(status, "000000"))


def build_report(
    summary: YearSummary, flags: pd.DataFrame, values: pd.DataFrame, thr: Thresholds, mode: str
) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Сводка"
    ws.append([f"Пульс инициатив — сводка за {summary.year} год"])
    ws["A1"].font = Font(bold=True, size=14, color=BLUE)
    note = (f"Данные на месяц {summary.as_of_month}; прогноз на конец года — линейная экстраполяция"
            if summary.is_current else "Закрытый год: прогноз = факт")
    if mode == "demo":
        note += ". ВНИМАНИЕ: синтетические демо-данные"
    ws.append([note])
    ws.append([])
    _header(ws, ["Место", "Инициатива", "Медиана % выполнения", "Медиана прогноза, % плана", "Статус",
                 "Показателей", "Есть критичный показатель", "Пояснение"], [8, 52, 16, 18, 16, 12, 16, 50])
    ranked = sorted(summary.initiatives, key=lambda i: -(i.median_fc_pct if i.median_fc_pct is not None else -1))
    for place, ini in enumerate(ranked, start=1):
        ws.append([place, ini.name, _round(ini.median_pct), _round(ini.median_fc_pct),
                   STATUS_LABELS[ini.status], ini.n_indicators, "да" if ini.has_crit else "", ini.note or ""])
        _paint_status(ws.cell(row=ws.max_row, column=5), ini.status)

    ws2 = wb.create_sheet("Детализация")
    _header(ws2, ["Инициатива", "Показатель", "Ед.", "Нарастающим итогом", "План", "Факт", "% выполнения",
                  "Прогноз на конец года", "Прогноз, % плана", "Статус", "Флаги"],
            [26, 70, 10, 12, 12, 12, 14, 16, 14, 16, 20])
    for r in summary.rows.itertuples():
        ws2.append([r.initiative, r.indicator, r.unit, "да" if r.is_cumulative else "", r.plan, r.fact,
                    _round(r.pct), _round(r.fc), _round(r.fc_pct), STATUS_LABELS[r.status], ", ".join(r.flags)])
        _paint_status(ws2.cell(row=ws2.max_row, column=10), r.status)
        ws2.cell(row=ws2.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    ws3 = wb.create_sheet("Флаги качества")
    _header(ws3, ["Флаг", "Что это", "Инициатива", "Показатель", "Год", "План", "Факт", "Пояснение"],
            [8, 26, 26, 60, 8, 12, 12, 60])
    for f in flags.itertuples():
        ws3.append([f.code, FLAG_TITLES.get(f.code, ""), f.initiative, f.indicator,
                    None if pd.isna(f.year) else int(f.year), f.plan, f.fact, f.detail])

    ws4 = wb.create_sheet("Данные")
    _header(ws4, ["Инициатива", "Показатель", "Ед.", "Нарастающим итогом", "Год", "План", "Факт",
                  "Данные на месяц", "Подтверждено", "Источник", "Загружено", "Комментарий"],
            [26, 70, 10, 12, 8, 12, 12, 12, 12, 14, 18, 40])
    for r in values.itertuples():
        ws4.append([r.initiative, r.indicator, r.unit, "да" if r.is_cumulative else "", int(r.year), r.plan,
                    r.fact, int(r.as_of_month), "да" if r.confirmed else "", r.source, r.loaded_at, r.comment or ""])

    ws5 = wb.create_sheet("Параметры")
    params = [
        ("Отчётный год", summary.year),
        ("Данные на месяц", summary.as_of_month),
        ("Порог «критично», %", f"< {thr.crit:g}"),
        ("Порог «риск», %", f"{thr.crit:g} – {thr.risk:g}"),
        ("«В норме», %", f"{thr.risk:g} – {thr.high:g}"),
        ("«План занижен», %", f"≥ {thr.high:g}"),
        ("Прогноз, обычный показатель", "F / M × 12"),
        ("Прогноз, нарастающим итогом", "F прошлого года + (F − F прошлого года) / M × 12"),
        ("Статус инициативы", "по медиане прогноза её показателей"),
        ("Режим данных", {"demo": "демо (синтетика)", "report": "отчёт Центра"}.get(mode, mode)),
        ("Сформировано", datetime.now().strftime("%d.%m.%Y %H:%M")),
    ]
    _header(ws5, ["Параметр", "Значение"], [34, 60])
    for k, v in params:
        ws5.append([k, v])

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row):
            for cell in row:
                if cell.value is not None and cell.row > 1:
                    cell.border = Border(bottom=THIN)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def protocol_xlsx(protocol: pd.DataFrame) -> bytes:
    """Протокол проверки файла для отправки куратору."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Протокол"
    _header(ws, list(protocol.columns), [16, 10, 10, 20, 60, 60])
    for row in protocol.itertuples(index=False):
        ws.append(list(row))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
