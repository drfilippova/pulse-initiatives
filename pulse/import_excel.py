"""Excel-шаблон для новых данных и чтение заполненных файлов (.xlsx / .csv)."""

from __future__ import annotations

from io import BytesIO, StringIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .catalog import INITIATIVES
from .validation import OPTIONAL, REQUIRED

MAX_FILE_MB = 5
HEADER_FILL = PatternFill("solid", fgColor="1E4FC2")
HEADER_FONT = Font(bold=True, color="FFFFFF")


class FileReadError(Exception):
    """Файл не удалось прочитать — сообщение готово для показа пользователю."""


def read_table(data: bytes, file_name: str) -> pd.DataFrame:
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise FileReadError(f"Файл больше {MAX_FILE_MB} МБ. Уберите лишние листы или разделите файл.")
    name = file_name.lower()
    try:
        if name.endswith(".csv"):
            text = None
            for enc in ("utf-8-sig", "cp1251"):
                try:
                    text = data.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                raise FileReadError("Не удалось определить кодировку CSV. Сохраните файл в UTF-8.")
            sep = ";" if text.split("\n", 1)[0].count(";") >= text.split("\n", 1)[0].count(",") else ","
            df = pd.read_csv(StringIO(text), sep=sep, dtype=str)
        elif name.endswith(".xlsx"):
            xls = pd.ExcelFile(BytesIO(data), engine="openpyxl")
            sheet = "Данные" if "Данные" in xls.sheet_names else xls.sheet_names[0]
            df = pd.read_excel(xls, sheet_name=sheet, dtype=object)
        else:
            raise FileReadError("Поддерживаются только файлы .xlsx и .csv.")
    except FileReadError:
        raise
    except Exception as exc:  # повреждённый файл, не тот формат и т. п.
        raise FileReadError(f"Не удалось прочитать файл: {exc.__class__.__name__}. Проверьте, что он не повреждён.")
    df = df.dropna(how="all")
    if df.empty:
        raise FileReadError("В файле нет строк с данными.")
    return df.reset_index(drop=True)


def build_template(indicators: pd.DataFrame, current_year: int, current_month: int) -> bytes:
    """Шаблон: лист «Данные» с выпадающими списками, «Справочники» и «Инструкция»."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    headers = REQUIRED + OPTIONAL
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = [34, 60, 8, 12, 12, 16, 40]
    for col, w in zip("ABCDEFG", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    ref = wb.create_sheet("Справочники")
    ref.append(["Инициатива", "Показатель", "Единица", "Нарастающим итогом"])
    for cell in ref[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    inis = [i.short for i in INITIATIVES]
    for r in indicators.itertuples():
        ref.append([r.initiative, r.name, r.unit, "да" if r.is_cumulative else "нет"])
    for col, w in zip("ABCD", [30, 80, 12, 20]):
        ref.column_dimensions[col].width = w
    ref["F1"] = "Список инициатив"
    ref["F1"].font = Font(bold=True)
    for i, short in enumerate(inis, start=2):
        ref[f"F{i}"] = short
    ref.column_dimensions["F"].width = 30
    n_ind = max(len(indicators) + 1, 2)

    last = 500
    dv_ini = DataValidation(type="list", formula1=f"=Справочники!$F$2:$F${len(inis) + 1}", allow_blank=False)
    dv_ind = DataValidation(type="list", formula1=f"=Справочники!$B$2:$B${n_ind}", allow_blank=False)
    dv_year = DataValidation(type="whole", operator="between", formula1="2022", formula2="2031")
    dv_month = DataValidation(type="whole", operator="between", formula1="1", formula2="12")
    dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0")
    for dv, rng in [(dv_ini, f"A2:A{last}"), (dv_ind, f"B2:B{last}"), (dv_year, f"C2:C{last}"),
                    (dv_num, f"D2:E{last}"), (dv_month, f"F2:F{last}")]:
        dv.error = "Значение не подходит — см. лист «Инструкция»"
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(rng)

    ins = wb.create_sheet("Инструкция")
    lines = [
        "Как заполнить шаблон «Пульс инициатив»",
        "",
        "1. Одна строка — один показатель одной инициативы за один год.",
        "2. Инициативу и показатель выбирайте из выпадающих списков (лист «Справочники»).",
        "3. План и Факт — числа без текста. Дробную часть можно писать через запятую.",
        "4. «Данные на месяц» — номер месяца, на который дан факт: для прошедшего года — 12.",
        "5. Не переименовывайте и не удаляйте колонки — иначе файл не пройдёт проверку.",
        "6. Единица измерения и признак «нарастающим итогом» берутся из справочника автоматически.",
    ]
    for line in lines:
        ins.append([line])
    ins["A1"].font = Font(bold=True, size=14, color="1E4FC2")
    ins.column_dimensions["A"].width = 100

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
